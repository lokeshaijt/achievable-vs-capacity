"""
Blend & Brand changeover
========================
Automates the manual Excel process that was used to count changeovers from the
Production Register.

Manual process (this module reproduces it step by step)
-------------------------------------------------------
 1. Production Register  -> keep rows with Transaction Type = "Production"      (Sheet1)
 2. Unique  Process_Date / Shift / Work Center / Machine No / Item Name          (Sheet2)
 3. Add the Blend name by VLOOKUP of the Item Name in "FG Blend Names"           (Sheet3)
 4. Unique  Process_Date / Shift / Work Center / Machine No / Blend name         (Sheet5, Sheet6)
 5. For every Process_Date / Shift / Work Center / Machine No:
        Brand changeover = distinct items  - 1     ("Brand Changeover"  pivot)
        Blend changeover = distinct blends - 1     ("blend change over" pivot)
 6. Sum by machine line and week  -> columns in the Achievable Vs Produced sheet

What gets written into the report workbook
------------------------------------------
 Changeover Detail  = steps 1-4  (unique item rows, blend looked up, "new blend in shift" flag)
 Changeover Calc    = step 5     (one row per machine-shift, live COUNTIFS / SUMIFS)
 CO Routing         = work center -> report machine line (editable)
 FG Blend Map       = FG item -> blend (first match, exactly like the VLOOKUP)
 + two column blocks (Blend / Brand changeover, week-wise) in the report sheet, and one
   all-machines total row under them
"""

import datetime
import io
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

NOT_IN_REPORT = "NOT IN REPORT"
UNKNOWN_PREFIX = "NOT IN FG LIST: "
PRODUCTION_TYPE = "Production"

SHEET_CALC = "Changeover Calc"
SHEET_DETAIL = "Changeover Detail"
SHEET_ROUTING = "CO Routing"
SHEET_FG = "FG Blend Map"

FIRST_DATA_ROW = 5          # header is row 4 on every changeover sheet


# ─────────────────────────────────────────────────────────────────────────────
# Reading the inputs
# ─────────────────────────────────────────────────────────────────────────────

def _open(src):
    if isinstance(src, (bytes, bytearray)):
        src = io.BytesIO(src)
    return openpyxl.load_workbook(src, data_only=True)


def _to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).date()
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_production_rows(src, sheet=None):
    """Steps 1-2: unique (date, shift, work center, machine, item) of 'Production' rows."""
    wb = _open(src)
    if sheet:
        ws = wb[sheet]
    elif "Default" in wb.sheetnames:
        ws = wb["Default"]
    elif "Production Register" in wb.sheetnames:
        ws = wb["Production Register"]
    else:
        ws = wb[wb.sheetnames[0]]

    hdr = next((r for r in range(1, 25) if str(ws.cell(r, 1).value).strip() == "Process_Date"), None)
    if hdr is None:
        raise ValueError("Production Register: could not find the header row (Process_Date).")
    cols = {str(ws.cell(hdr, c).value).strip(): c
            for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value is not None}
    need = ["Process_Date", "Shift", "Work Center", "Machine No", "Item Name", "Transaction Type"]
    missing = [n for n in need if n not in cols]
    if missing:
        raise ValueError(f"Production Register is missing column(s): {', '.join(missing)}")

    seen = OrderedDict()
    for r in range(hdr + 1, ws.max_row + 1):
        d = _to_date(ws.cell(r, cols["Process_Date"]).value)
        if d is None:
            continue
        if str(ws.cell(r, cols["Transaction Type"]).value).strip() != PRODUCTION_TYPE:
            continue
        key = (d,
               _txt(ws.cell(r, cols["Shift"]).value),
               _txt(ws.cell(r, cols["Work Center"]).value),
               _txt(ws.cell(r, cols["Machine No"]).value),
               _txt(ws.cell(r, cols["Item Name"]).value))
        seen[key] = 1
    return sorted(seen)


def read_fg_blend_map(src):
    """Step 3 lookup table. First match per FG name wins (same as the VLOOKUP).
    Returns OrderedDict  UPPER(fg name) -> (fg name, blend name)."""
    wb = _open(src)
    ws = wb[wb.sheetnames[0]]
    hdr = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)
           if ws.cell(1, c).value is not None}
    if "FG Name" not in hdr or "Blend Name" not in hdr:
        raise ValueError("FG Blend Names file needs the columns 'FG Name' and 'Blend Name'.")
    out = OrderedDict()
    for r in range(2, ws.max_row + 1):
        fg = ws.cell(r, hdr["FG Name"]).value
        bl = ws.cell(r, hdr["Blend Name"]).value
        if fg is None or bl is None or str(fg).strip() == "":
            continue
        out.setdefault(str(fg).upper(), (str(fg), str(bl)))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# The model (pure Python, also used to cross-check the Excel formulas)
# ─────────────────────────────────────────────────────────────────────────────

class ChangeoverModel:
    def __init__(self, detail, calc, routing_rows, fg_rows, unknown_items, line_of_wc):
        self.detail = detail                # (date, shift, wc, machine, item, blend)
        self.calc = calc                    # dicts: date, shift, wc, machine, items, blends, line
        self.routing_rows = routing_rows    # (work center, report line)
        self.fg_rows = fg_rows              # (fg name, blend)
        self.unknown_items = unknown_items  # item names with no blend in the FG list
        self.line_of_wc = line_of_wc

    @property
    def totals(self):
        return (sum(c["blends"] - 1 for c in self.calc), sum(c["items"] - 1 for c in self.calc))

    def by_line_week(self, weekno):
        """{(line, week): (blend_co, brand_co)} with week from weekno(date)."""
        out = defaultdict(lambda: [0, 0])
        for c in self.calc:
            k = (c["line"], weekno(c["date"]))
            out[k][0] += c["blends"] - 1
            out[k][1] += c["items"] - 1
        return {k: tuple(v) for k, v in out.items()}

    def not_in_report(self):
        """{work center: (blend_co, brand_co)} for work centers with no report line."""
        out = defaultdict(lambda: [0, 0])
        for c in self.calc:
            if c["line"] == NOT_IN_REPORT:
                out[c["wc"]][0] += c["blends"] - 1
                out[c["wc"]][1] += c["items"] - 1
        return {k: tuple(v) for k, v in out.items()}


def build_model(pr_src, fg_src, routing_map, report_lines):
    rows = read_production_rows(pr_src)
    fg = read_fg_blend_map(fg_src)
    report_lines = set(report_lines)

    detail, unknown = [], set()
    for d, s, wc, m, item in rows:
        hit = fg.get(item.upper())
        if hit:
            blend = hit[1]
        else:
            blend = UNKNOWN_PREFIX + item       # unknown item = its own blend (never hides a changeover)
            unknown.add(item)
        detail.append((d, s, wc, m, item, blend))

    wcs = sorted({r[2] for r in rows})
    line_of_wc = {}
    for wc in wcs:
        g = routing_map.get(wc)
        line_of_wc[wc] = g if g in report_lines else NOT_IN_REPORT

    items, blends = defaultdict(set), defaultdict(set)
    for d, s, wc, m, item, blend in detail:
        items[(d, s, wc, m)].add(item)
        blends[(d, s, wc, m)].add(blend)
    calc = [dict(date=k[0], shift=k[1], wc=k[2], machine=k[3],
                 items=len(items[k]), blends=len(blends[k]), line=line_of_wc[k[2]])
            for k in sorted(items)]

    return ChangeoverModel(detail, calc,
                           [(wc, line_of_wc[wc]) for wc in wcs],
                           list(fg.values()), sorted(unknown), line_of_wc)


# ─────────────────────────────────────────────────────────────────────────────
# Writing the supporting sheets
# ─────────────────────────────────────────────────────────────────────────────

_FN = "Calibri"
_THIN = Side(style="thin", color="FFB0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)            # light grey (table headers, total row)
_BLACK_SIDE = Side(style="thin")                                               # automatic (black) colour
_BLACK = Border(left=_BLACK_SIDE, right=_BLACK_SIDE, top=_BLACK_SIDE, bottom=_BLACK_SIDE)   # report body
_TITLE_FILL = PatternFill(start_color="FFE8D5B0", end_color="FFE8D5B0", fill_type="solid")
_HDR_FILL = PatternFill(start_color="FF4A3728", end_color="FF4A3728", fill_type="solid")
_INPUT_FILL = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
_TOT_FILL = PatternFill(start_color="FFCDC4A0", end_color="FFCDC4A0", fill_type="solid")
_WARN_FILL = PatternFill(start_color="FFFFE4E1", end_color="FFFFE4E1", fill_type="solid")
_NORM = Font(name=_FN, size=10)
_BOLD = Font(name=_FN, bold=True, size=10)
_TITLE = Font(name=_FN, bold=True, size=13, color="FF8B1A1A")
_NOTE = Font(name=_FN, size=9, italic=True, color="FF444444")
_HDR = Font(name=_FN, bold=True, size=10, color="FFFFFFFF")
_DATE_FMT = "d-mmm-yy"


def _title(ws, text, note, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, text).font = _TITLE
    ws.cell(1, 1).fill = _TITLE_FILL
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, note).font = _NOTE
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 40


def _header(ws, labels, widths):
    ws.row_dimensions[4].height = 32
    for c, (h, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(4, c, h)
        cell.font = _HDR
        cell.fill = _HDR_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(c)].width = w


DEFAULT_WEEK1_MONDAY = datetime.date(2025, 12, 29)     # Monday of the week holding 1 Jan 2026 = week 1


def _week_formula(cell, week1_monday):
    """Continuous week number (never restarts): weeks counted from week1_monday, like the report headers."""
    return (f"=INT(({cell}-DATE({week1_monday.year},{week1_monday.month},{week1_monday.day}))/7)+1")


def write_changeover_sheets(wb, model, report_lines, week1_monday=DEFAULT_WEEK1_MONDAY):
    """Adds the four supporting sheets. Returns the ranges the report columns refer to."""
    n_det, n_calc = len(model.detail), len(model.calc)
    n_fg, n_rt = len(model.fg_rows), len(model.routing_rows)
    last_det = FIRST_DATA_ROW + n_det - 1
    last_calc = FIRST_DATA_ROW + n_calc - 1
    last_fg = FIRST_DATA_ROW + n_fg - 1
    last_rt = FIRST_DATA_ROW + n_rt - 1

    # ── FG Blend Map ─────────────────────────────────────────────────────
    ws = wb.create_sheet(SHEET_FG)
    _title(ws, "FG Blend Map — FG item → Blend name",
           "Source: FG Blend Names file. One row per FG name; where an FG has several blends the FIRST one is used "
           "(same as the VLOOKUP used in the manual process).", 2)
    _header(ws, ["FG Name", "Blend Name"], [60, 46])
    for i, (fg, bl) in enumerate(model.fg_rows):
        r = FIRST_DATA_ROW + i
        ws.cell(r, 1, fg).font = _NORM
        ws.cell(r, 2, bl).font = _NORM
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    # ── CO Routing ───────────────────────────────────────────────────────
    ws = wb.create_sheet(SHEET_ROUTING)
    _title(ws, "CO Routing — Work Center → machine line in the report",
           "Edit column B (yellow) to move a work center to a different machine line. Use a Machine Line name exactly as in "
           f"the report, or '{NOT_IN_REPORT}'. The report columns and Changeover Calc update automatically.", 4)
    _header(ws, ["Work Center", "Report Line", "", "Machine lines in report (list for dropdown)"], [34, 30, 3, 34])
    for i, (wc, line) in enumerate(model.routing_rows):
        r = FIRST_DATA_ROW + i
        ws.cell(r, 1, wc).font = _NORM
        c = ws.cell(r, 2, line)
        c.font = _NORM
        c.fill = _INPUT_FILL
        if line == NOT_IN_REPORT:
            c.font = Font(name=_FN, size=10, italic=True, color="FF990000")
    choices = list(report_lines) + [NOT_IN_REPORT]
    for i, name in enumerate(choices):
        ws.cell(FIRST_DATA_ROW + i, 4, name).font = _NORM
    last_choice = FIRST_DATA_ROW + len(choices) - 1
    dv = DataValidation(type="list", formula1=f"=$D${FIRST_DATA_ROW}:$D${last_choice}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B{FIRST_DATA_ROW}:B{last_rt}")
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    # ── Changeover Detail (steps 1-4) ───────────────────────────────────
    ws = wb.create_sheet(SHEET_DETAIL)
    _title(ws, "Changeover Detail — unique Date / Shift / Work Center / Machine / Item, with Blend",
           "Steps 1-4 of the manual process: Production rows only → unique item per machine-shift → Blend by lookup in "
           "FG Blend Map. Column I = 1 the first time a blend appears in a machine-shift (so SUM = distinct blends). "
           f"Items missing from the FG list show '{UNKNOWN_PREFIX}…' and are counted as their own blend.", 9)
    ws.cell(3, 6, "Items missing from FG Blend Map:").font = _BOLD
    ws.cell(3, 6).alignment = Alignment(horizontal="right")
    ws.cell(3, 7, f'=COUNTIF($G${FIRST_DATA_ROW}:$G${last_det},"{UNKNOWN_PREFIX}*")').font = _BOLD
    _header(ws, ["Process Date", "Week", "Shift", "Work Center", "Machine No", "Item Name (FG)",
                 "Blend Name", "Key (date|shift|work center|machine)", "New blend in\nmachine-shift?"],
            [12, 7, 7, 30, 12, 58, 44, 46, 13])
    for i, (d, s, wc, m, item, blend) in enumerate(model.detail):
        r = FIRST_DATA_ROW + i
        ws.cell(r, 1, d).number_format = _DATE_FMT
        ws.cell(r, 2, _week_formula(f"A{r}", week1_monday))
        ws.cell(r, 3, s)
        ws.cell(r, 4, wc)
        ws.cell(r, 5, m)
        ws.cell(r, 6, item)
        ws.cell(r, 7, f"=IFERROR(VLOOKUP(F{r},'{SHEET_FG}'!$A${FIRST_DATA_ROW}:$B${last_fg},2,0),"
                      f"\"{UNKNOWN_PREFIX}\"&F{r})")
        ws.cell(r, 8, f'=TEXT(A{r},"yyyymmdd")&"|"&C{r}&"|"&D{r}&"|"&E{r}')
        ws.cell(r, 9, f"=IF(COUNTIFS($H${FIRST_DATA_ROW}:H{r},H{r},$G${FIRST_DATA_ROW}:G{r},G{r})=1,1,0)")
        for c in range(1, 10):
            ws.cell(r, c).font = _NORM
    ws.conditional_formatting.add(
        f"A{FIRST_DATA_ROW}:I{last_det}",
        FormulaRule(formula=[f'LEFT($G{FIRST_DATA_ROW},{len(UNKNOWN_PREFIX)})="{UNKNOWN_PREFIX}"'], fill=_WARN_FILL))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{last_det}"
    ws.sheet_view.showGridLines = False

    # ── Changeover Calc (step 5) ─────────────────────────────────────────
    ws = wb.create_sheet(SHEET_CALC)
    _title(ws, "Changeover Calc — one row per Date / Shift / Work Center / Machine",
           "Step 5 of the manual process (the two pivots). Brand changeover = distinct FG items − 1; Blend changeover = "
           "distinct blends − 1. Week = weeks counted from Mon 29-Dec-2025 (same numbering as the report headers). "
           "Machine Line comes from 'CO Routing'. The report adds these up by machine line and week.", 11)
    ws.cell(3, 9, "TOTAL (all machines) →").font = _BOLD
    ws.cell(3, 9).alignment = Alignment(horizontal="right")
    ws.cell(3, 10, f"=SUM(J{FIRST_DATA_ROW}:J{last_calc})").font = _BOLD
    ws.cell(3, 11, f"=SUM(K{FIRST_DATA_ROW}:K{last_calc})").font = _BOLD
    _header(ws, ["Process Date", "Week", "Shift", "Work Center", "Machine No", "Machine Line\n(report)",
                 "Key", "Distinct\nitems", "Distinct\nblends", "Brand\nchangeover\n(items − 1)",
                 "Blend\nchangeover\n(blends − 1)"],
            [12, 7, 7, 30, 12, 28, 46, 10, 10, 13, 13])
    ws.row_dimensions[4].height = 44
    det_key = f"'{SHEET_DETAIL}'!$H${FIRST_DATA_ROW}:$H${last_det}"
    det_flag = f"'{SHEET_DETAIL}'!$I${FIRST_DATA_ROW}:$I${last_det}"
    for i, c in enumerate(model.calc):
        r = FIRST_DATA_ROW + i
        ws.cell(r, 1, c["date"]).number_format = _DATE_FMT
        ws.cell(r, 2, _week_formula(f"A{r}", week1_monday))
        ws.cell(r, 3, c["shift"])
        ws.cell(r, 4, c["wc"])
        ws.cell(r, 5, c["machine"])
        ws.cell(r, 6, f"=IFERROR(VLOOKUP(D{r},'{SHEET_ROUTING}'!$A${FIRST_DATA_ROW}:$B${last_rt},2,0),"
                      f"\"{NOT_IN_REPORT}\")")
        ws.cell(r, 7, f'=TEXT(A{r},"yyyymmdd")&"|"&C{r}&"|"&D{r}&"|"&E{r}')
        ws.cell(r, 8, f"=COUNTIFS({det_key},G{r})")
        ws.cell(r, 9, f"=SUMIFS({det_flag},{det_key},G{r})")
        ws.cell(r, 10, f"=H{r}-1")
        ws.cell(r, 11, f"=I{r}-1")
        for col in range(1, 12):
            ws.cell(r, col).font = _NORM
    grey = Font(name=_FN, size=10, italic=True, color="FF808080")
    ws.conditional_formatting.add(
        f"A{FIRST_DATA_ROW}:K{last_calc}",
        FormulaRule(formula=[f'$F{FIRST_DATA_ROW}="{NOT_IN_REPORT}"'], font=grey))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last_calc}"
    ws.sheet_view.showGridLines = False

    # keep a predictable tab order: Calc, Detail, Routing, FG map (after whatever is already there)
    names = [SHEET_CALC, SHEET_DETAIL, SHEET_ROUTING, SHEET_FG]
    wb._sheets = [sh for sh in wb._sheets if sh.title not in names] + [wb[n] for n in names]

    return dict(calc_week=f"'{SHEET_CALC}'!$B${FIRST_DATA_ROW}:$B${last_calc}",
                calc_line=f"'{SHEET_CALC}'!$F${FIRST_DATA_ROW}:$F${last_calc}",
                calc_brand=f"'{SHEET_CALC}'!$J${FIRST_DATA_ROW}:$J${last_calc}",
                calc_blend=f"'{SHEET_CALC}'!$K${FIRST_DATA_ROW}:$K${last_calc}",
                last_calc=last_calc, last_detail=last_det)


# ─────────────────────────────────────────────────────────────────────────────
# Writing the two column blocks into the Achievable Vs Produced sheet
# ─────────────────────────────────────────────────────────────────────────────

_BLEND_HDR1 = PatternFill(start_color="FFC55A11", end_color="FFC55A11", fill_type="solid")
_BLEND_HDR2 = PatternFill(start_color="FFF4B183", end_color="FFF4B183", fill_type="solid")
_BRAND_HDR1 = PatternFill(start_color="FF1F7A8C", end_color="FF1F7A8C", fill_type="solid")
_BRAND_HDR2 = PatternFill(start_color="FFB7DEE8", end_color="FFB7DEE8", fill_type="solid")
_FOOT_FILL = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
_GREEN = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")   # Brand block body
_CO_FMT = '#,##0;\\-#,##0;\\-'


def write_changeover_columns(ws, refs, first_col, weeks, line_rows, total_groups, footer_row=None,
                             line_col="C"):
    """
    first_col    : first column of the Blend block (Brand block follows it)
    weeks        : e.g. [36, 37, 38, 39]
    line_rows    : {sheet row: machine line name}
    total_groups : {TOTAL row: (first row, last row)} of the lines it adds up
    footer_row   : row for the all-machines changeover total (None = none). It has no label and
                   covers every machine, including work centers that have no line in the report.
    line_col     : column letter of the report's "Machine Line" column
    """
    n = len(weeks)
    blocks = [("Blend Changeover", first_col, refs["calc_blend"], _BLEND_HDR1, _BLEND_HDR2, None),
              ("Brand Changeover", first_col + n, refs["calc_brand"], _BRAND_HDR1, _BRAND_HDR2, _GREEN)]
    white_b = Font(name=_FN, bold=True, size=12, color="FFFFFFFF")
    dark_b = Font(name=_FN, bold=True, size=10, color="FF000000")

    def week_expr(col):
        return f'VALUE(SUBSTITUTE({L(col)}$2,"W #",""))'

    for title, c0, rng, fill1, fill2, body_fill in blocks:
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + n - 1)
        h = ws.cell(1, c0, title)
        h.font, h.fill = white_b, fill1
        h.alignment = Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(weeks):
            col = c0 + i
            cell = ws.cell(2, col, f"W #{w}")
            cell.font, cell.fill, cell.border = dark_b, fill2, _BLACK
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in line_rows:
            for i in range(n):
                col = c0 + i
                cell = ws.cell(r, col,
                               f"=SUMIFS({rng},{refs['calc_line']},${line_col}{r},{refs['calc_week']},{week_expr(col)})")
                cell.font, cell.border, cell.number_format = _NORM, _BLACK, _CO_FMT
                cell.alignment = Alignment(horizontal="right")
                if body_fill is not None:
                    cell.fill = body_fill
        for r, (s_, e_) in total_groups.items():
            for i in range(n):
                col = c0 + i
                cell = ws.cell(r, col, f"=SUM({L(col)}{s_}:{L(col)}{e_})")
                cell.font, cell.fill, cell.border = _BOLD, (body_fill or _TOT_FILL), _BLACK
                cell.number_format = _CO_FMT
                cell.alignment = Alignment(horizontal="right")

        if footer_row:
            for i in range(n):
                col = c0 + i
                cell = ws.cell(footer_row, col, f"=SUMIFS({rng},{refs['calc_week']},{week_expr(col)})")
                cell.font, cell.fill, cell.border = _NORM, _FOOT_FILL, _BORDER
                cell.number_format = _CO_FMT
                cell.alignment = Alignment(horizontal="right")
    if footer_row:
        ws.row_dimensions[footer_row].height = 18


def add_changeover(wb, ws, pr_src, fg_src, routing_map, report_lines, weeks, first_col,
                   line_rows, total_groups, footer_row=None, week1_monday=DEFAULT_WEEK1_MONDAY,
                   line_col="C"):
    """One call: build the model, write the supporting sheets and the two column blocks."""
    model = build_model(pr_src, fg_src, routing_map, report_lines)
    refs = write_changeover_sheets(wb, model, report_lines, week1_monday)
    write_changeover_columns(ws, refs, first_col, weeks, line_rows, total_groups, footer_row, line_col)
    return model
