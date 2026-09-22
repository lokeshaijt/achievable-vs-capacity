"""
Achievable Vs Produced — Report Generator
Converts Production Register → container quantities → single-sheet report
"""

import datetime
import io
import re
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color
from openpyxl.utils import get_column_letter


# ── Layout: OVW rows 3-38 (region, sub, line, no_machines, no_shifts) ────────
# TOTAL rows included so we can reproduce the same grouping structure
LAYOUT = [
    ("AFRICA",    "DC",         "C250 AF",                  2,   2),
    (None,        None,         "CONSTANTA ENV",            32,  2),
    (None,        None,         "PERFECTA ENV",             3,   2),
    (None,        None,         "MD20 ENV",                 4,   2),
    ("__TOTAL__", None,         None,                       None,None),
    (None,        "SC",         "MAISA TAGLESS",            7,   1),
    ("__TOTAL__", None,         None,                       None,None),
    (None,        "PREMIX",     "INSTANT COFFEE STICK PACK",1,   1),
    (None,        None,         "Instant Tea",              1,   1),
    (None,        None,         "PEARL PACK-PREMIX",        2,   1),
    (None,        None,         "VIKING-MULTITRACK",        1,   1),
    ("__TOTAL__", None,         None,                       None,None),
    ("AUSTRALIA", "DC TAG",     "CONSTANTA TAG (D)",        5,   1),
    ("__TOTAL__", None,         None,                       None,None),
    ("EUROPE",    "DC SF CRIMP","C250",                     6,   2),
    (None,        None,         "VARIETIES PACK",           1,   2),
    ("__TOTAL__", None,         None,                       None,None),
    (None,        "DC TAG",     "PERFECTA  K 45 TAG",       1,   2),
    (None,        None,         "MD20 4GM TAG",             1,   1),
    ("__TOTAL__", None,         None,                       None,None),
    (None,        "SC",         "UNIVERSAL POT BAG",        1,   1),
    (None,        None,         "UNIVERSAL TWIN BAG",       1,   1),
    ("__TOTAL__", None,         None,                       None,None),
    (None,        "MISC",       "C250F3 HS",                1,   0),
    (None,        None,         "MD20 HS",                  2,   2),
    ("__TOTAL__", None,         None,                       None,None),
    ("RUSSIA",    "DC TAG",     "CONSTANTA TAG",            3,   1),
    ("__TOTAL__", None,         None,                       None,None),
    ("USA",       "TREE HOUSE", "BREW MAGIC",               1,   1),
    (None,        None,         "C250-A",                   3,   2),
    (None,        None,         "MAISA ENV",                2,   1),
    (None,        None,         "PERFACTA TAG",             2,   2),
    (None,        None,         "CONSTANTA FAMILY SIZE",    2,   2),
    (None,        "HF Co",      "Pearl Pack-FFS POUCH",     1.5, 1),
    (None,        None,         "C250F3 ENV",               1,   0),
    ("__TOTAL__", None,         None,                       None,None),
]

# ── MC MASTER factors ─────────────────────────────────────────────────────────
MC_FACTORS = {
    "BREW MAGIC":               {"rate":8,   "shiftmin":540, "tbgs_cfc":32,   "cfc_cntr":3600, "eff":0.9},
    "C250":                     {"rate":250, "shiftmin":524, "tbgs_cfc":480,  "cfc_cntr":3600, "eff":0.8},
    "C250 AF":                  {"rate":250, "shiftmin":524, "tbgs_cfc":1200, "cfc_cntr":3340, "eff":0.9},
    "C250-A":                   {"rate":250, "shiftmin":524, "tbgs_cfc":480,  "cfc_cntr":3600, "eff":0.9},
    "C250F3 ENV":               {"rate":250, "shiftmin":520, "tbgs_cfc":240,  "cfc_cntr":7200, "eff":0.9},
    "C250F3 HS":                {"rate":250, "shiftmin":522, "tbgs_cfc":240,  "cfc_cntr":7200, "eff":0.9},
    "CONSTANTA ENV":            {"rate":125, "shiftmin":553, "tbgs_cfc":1125, "cfc_cntr":3340, "eff":0.9},
    "CONSTANTA FAMILY SIZE":    {"rate":95,  "shiftmin":540, "tbgs_cfc":144,  "cfc_cntr":5720, "eff":0.9},
    "CONSTANTA TAG":            {"rate":120, "shiftmin":543, "tbgs_cfc":1200, "cfc_cntr":2688, "eff":0.9},
    "CONSTANTA TAG (D)":        {"rate":120, "shiftmin":543, "tbgs_cfc":1200, "cfc_cntr":2688, "eff":0.9},
    "INSTANT COFFEE STICK PACK":{"rate":40,  "shiftmin":540, "tbgs_cfc":1350, "cfc_cntr":3000, "eff":0.9},
    "Instant Tea":              {"rate":14,  "shiftmin":550, "tbgs_cfc":60,   "cfc_cntr":15120,"eff":0.9},
    "MAISA ENV":                {"rate":100, "shiftmin":544, "tbgs_cfc":1200, "cfc_cntr":1656, "eff":0.9},
    "MAISA TAGLESS":            {"rate":120, "shiftmin":555, "tbgs_cfc":2400, "cfc_cntr":2800, "eff":0.9},
    "MD20 4GM TAG":             {"rate":125, "shiftmin":548, "tbgs_cfc":480,  "cfc_cntr":4032, "eff":0.9},
    "MD20 ENV":                 {"rate":185, "shiftmin":528, "tbgs_cfc":1200, "cfc_cntr":3340, "eff":0.9},
    "MD20 HS":                  {"rate":170, "shiftmin":526, "tbgs_cfc":300,  "cfc_cntr":6500, "eff":0.9},
    "Pearl Pack-FFS POUCH":     {"rate":24,  "shiftmin":550, "tbgs_cfc":24,   "cfc_cntr":1200, "eff":0.9},
    "PEARL PACK-PREMIX":        {"rate":30,  "shiftmin":550, "tbgs_cfc":120,  "cfc_cntr":6000, "eff":0.9},
    "PERFACTA TAG":             {"rate":300, "shiftmin":525, "tbgs_cfc":600,  "cfc_cntr":4620, "eff":0.9},
    "PERFECTA  K 45 TAG":       {"rate":250, "shiftmin":525, "tbgs_cfc":600,  "cfc_cntr":5040, "eff":0.9},
    "PERFECTA ENV":             {"rate":280, "shiftmin":539, "tbgs_cfc":1200, "cfc_cntr":3340, "eff":0.9},
    "UNIVERSAL POT BAG":        {"rate":272, "shiftmin":540, "tbgs_cfc":960,  "cfc_cntr":3220, "eff":0.9},
    "UNIVERSAL TWIN BAG":       {"rate":420, "shiftmin":541, "tbgs_cfc":960,  "cfc_cntr":3220, "eff":0.9},
    "VARIETIES PACK":           {"rate":250, "shiftmin":524, "tbgs_cfc":480,  "cfc_cntr":3600, "eff":0.9},
    "VIKING-MULTITRACK":        {"rate":56,  "shiftmin":550, "tbgs_cfc":60,   "cfc_cntr":15120,"eff":0.9},
}

# ── Routing map: work-center name → capacity group (machine line) ─────────────
ROUTING_MAP = {
    "BREW MAGIC":           "BREW MAGIC",
    "C250":                 "C250",
    "C250 AF":              "C250 AF",
    "C250-A":               "C250-A",
    "C250F3 ENV":           "C250F3 ENV",
    "C250F3 HS":            "C250F3 HS",
    "CONSTANTA ENV":        "CONSTANTA ENV",
    "CONSTANTA ENV BT":     "CONSTANTA ENV",
    "CONSTANTA ENV BT 25":  "CONSTANTA ENV",
    "CONSTANTA ENV F":      "CONSTANTA ENV",
    "CONSTANTA ENV G":      "CONSTANTA ENV",
    "CONSTANTA ENV H":      "CONSTANTA ENV",
    "CONSTANTA ENV HF":     "CONSTANTA ENV",
    "CONSTANTA ENV M":      "CONSTANTA ENV",
    "CONSTANTA FAMILY SIZE":"CONSTANTA FAMILY SIZE",
    "CONSTANTA TAG":        "CONSTANTA TAG",
    "CONSTANTA TAG (D)":    "CONSTANTA TAG (D)",
    "INSTANT COFFEE STICK PACK":"INSTANT COFFEE STICK PACK",
    "Instant Tea":          "Instant Tea",
    "MAISA ENV":            "MAISA ENV",
    "MAISA TAGLESS":        "MAISA TAGLESS",
    "MD20 4GM TAG":         "MD20 4GM TAG",
    "MD20 ENV":             "MD20 ENV",
    "MD20 HS":              "MD20 HS",
    "PEARL PACK-NANO":      "PEARL PACK-NANO",
    "PEARL PACK-PREMIX":    "PEARL PACK-PREMIX",
    "PEARL PACK-PREMIX 2":  "PEARL PACK-PREMIX",
    "PERFACTA TAG":         "PERFACTA TAG",
    "PERFECTA  K 45 TAG":   "PERFECTA  K 45 TAG",
    "PERFECTA ENV":         "PERFECTA ENV",
    "Pearl Pack-FFS POUCH": "Pearl Pack-FFS POUCH",
    "UNIVERSAL POT BAG":    "UNIVERSAL POT BAG",
    "UNIVERSAL TWIN BAG":   "UNIVERSAL TWIN BAG",
    "VARIETIES PACK":       "VARIETIES PACK",
    "VIKING-MULTITRACK":    "VIKING-MULTITRACK",
}

# Work centers to skip (non-packing lines)
SKIP_WORK_CENTERS = {
    "CFC FORMING", "RECLAIM", "MANUAL", "UNFOLD", "BAG FORMING",
    "MISC", "PAKONA 250GMS/500GMS/1000GMS", "FFS POUCH", "FFS POUCH - 13G- 40G",
}
SKIP_UOMS = {"KGS", "BAG", "POUCH", "ENV", "PKT", "BOX", "CASE", "PCS"}


# ── Production Register extraction ────────────────────────────────────────────

def extract_prod_register(pr_bytes):
    """Parse Production Register. Returns list of (date, group, item_name, qty, uom)."""
    wb = openpyxl.load_workbook(io.BytesIO(pr_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in range(8, ws.max_row + 1):
        pdate = ws.cell(r, 1).value
        wc    = ws.cell(r, 3).value
        item_name = ws.cell(r, 6).value
        qty   = ws.cell(r, 8).value
        uom   = ws.cell(r, 10).value
        if not pdate or not wc:
            continue
        if wc in SKIP_WORK_CENTERS:
            continue
        if uom in SKIP_UOMS:
            continue
        if uom not in ("CTN", "CFC"):
            continue
        group = ROUTING_MAP.get(wc)
        if not group:
            continue
        d = pdate.date() if isinstance(pdate, datetime.datetime) else pdate
        rows.append((d, group, item_name or "", qty or 0, uom))
    return rows


def load_tbgs_per_ctn(ref_bytes):
    """Load TBGS PER CTN lookup from reference workbook."""
    wb = openpyxl.load_workbook(io.BytesIO(ref_bytes), data_only=True)
    ws = wb["TBGS PER CTN"]
    table = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        val  = ws.cell(r, 3).value
        if name and val:
            table[name] = val
    return table


# ── Week numbering ────────────────────────────────────────────────────────────

_JAN1_2026    = datetime.date(2026, 1, 1)
_WEEK1_MONDAY = _JAN1_2026 - datetime.timedelta(days=_JAN1_2026.weekday())


def weeknum(d):
    monday = d - datetime.timedelta(days=d.weekday())
    return ((monday - _WEEK1_MONDAY).days // 7) + 1


def week_label(wn):
    return f"W #{wn}"


# ── Container conversion ───────────────────────────────────────────────────────

def prod_to_containers(prod_rows, tbgs_per_ctn, target_weeks):
    """Convert production rows to containers per (group, weeknum)."""
    result = defaultdict(float)
    skipped_names = set()

    for d, group, item_name, qty, uom in prod_rows:
        wn = weeknum(d)
        if wn not in target_weeks:
            continue
        factors = MC_FACTORS.get(group, {})
        tpc = factors.get("tbgs_cfc", 0)
        cpc = factors.get("cfc_cntr", 0)
        eff = factors.get("eff", 0.9)
        if not tpc or not cpc:
            continue
        if uom == "CTN":
            tbgs_ctn = tbgs_per_ctn.get(item_name)
            if not tbgs_ctn:
                # Derive from product name: "20 DC SF ENV" → 20
                m = re.search(r"\b(\d+)\s+DC\b", item_name or "")
                tbgs_ctn = int(m.group(1)) if m else None
            if not tbgs_ctn:
                skipped_names.add(item_name)
                continue
            cfc = qty * tbgs_ctn / tpc
        else:
            cfc = qty
        result[(group, wn)] += cfc / cpc

    return dict(result), skipped_names


def compute_capacity(line, machines, shifts):
    """Compute achievable weekly capacity in containers."""
    f = MC_FACTORS.get(line, {})
    rate     = f.get("rate", 0)
    shiftmin = f.get("shiftmin", 0)
    tbgs_cfc = f.get("tbgs_cfc", 0)
    cfc_cntr = f.get("cfc_cntr", 0)
    eff      = f.get("eff", 0.9)
    if not (rate and shiftmin and tbgs_cfc and cfc_cntr and machines and shifts):
        return 0.0
    return rate * shiftmin * machines * shifts * 6 * eff / tbgs_cfc / cfc_cntr


# ── Styles ────────────────────────────────────────────────────────────────────

def _styles():
    FN = "Calibri"
    thin = Side(style="thin", color="FFB0B0B0")
    return {
        "FN": FN,
        "BORDER": Border(left=thin, right=thin, top=thin, bottom=thin),
        "TITLE_L_FILL": PatternFill(start_color="FFE8D5B0", end_color="FFE8D5B0", fill_type="solid"),
        "TITLE_C_FILL": PatternFill(fgColor=Color(theme=5, tint=-0.5), fill_type="solid"),
        "TITLE_R_FILL": PatternFill(fgColor=Color(theme=7, tint=-0.5), fill_type="solid"),
        "HDR_FILL":     PatternFill(start_color="FFB8C4D6", end_color="FFB8C4D6", fill_type="solid"),
        "CNTR_HDR_FILL":PatternFill(fgColor=Color(theme=3, tint=-0.25), fill_type="solid"),
        "PCT_HDR_FILL": PatternFill(fgColor=Color(theme=6, tint=-0.5),  fill_type="solid"),
        "TOTAL_FILL":   PatternFill(start_color="FFCDC4A0", end_color="FFCDC4A0", fill_type="solid"),
        "REGION_FILL":  PatternFill(start_color="FFD6D6D6", end_color="FFD6D6D6", fill_type="solid"),
        "INPUT_FILL":   PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid"),
        "TITLE_L_FONT": Font(name=FN, bold=True, size=14, color="FF8B1A1A"),
        "TITLE_C_FONT": Font(name=FN, bold=True, size=12, color="FFFFFFFF"),
        "TITLE_R_FONT": Font(name=FN, bold=True, size=12, color="FFFFFFFF"),
        "HDR_FONT":     Font(name=FN, bold=True, size=10),
        "NORM":         Font(name=FN, size=10),
        "BOLD":         Font(name=FN, bold=True, size=10),
        "REGION_FONT":  Font(name=FN, bold=True, size=10),
        "OVER_FONT":    Font(name=FN, size=10, color="FFCC3300"),
        "OVER_FONT_B":  Font(name=FN, bold=True, size=10, color="FFCC3300"),
        "NUMFMT_CNTR":  '_ * #,##0.0_ ;_ * \\-#,##0.0_ ;_ * "-"??_ ;_ @_ ',
        "NUMFMT_PCT":   '0"%"',
    }


# ── Main report builder ───────────────────────────────────────────────────────

def generate_avp_report(pr_bytes, ref_bytes, target_weeks):
    """
    Parameters
    ----------
    pr_bytes      : bytes  Production Register xlsx
    ref_bytes     : bytes  Reference workbook (for TBGS PER CTN)
    target_weeks  : list   e.g. [35, 36, 37, 38]

    Returns
    -------
    (bytes, set)  — xlsx bytes, set of unmatched item names
    """
    # Load data
    tbgs_per_ctn = load_tbgs_per_ctn(ref_bytes)
    prod_rows    = extract_prod_register(pr_bytes)
    achieved, skipped = prod_to_containers(prod_rows, tbgs_per_ctn, set(target_weeks))

    week_labels = [week_label(wn) for wn in target_weeks]
    n_weeks     = len(target_weeks)
    CNTR_FIRST  = 7   # col G
    PCT_FIRST   = CNTR_FIRST + n_weeks  # col K (for n_weeks=4)
    LAST_COL    = PCT_FIRST + n_weeks - 1

    S = _styles()
    FN, BORDER = S["FN"], S["BORDER"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = " achievable vs produced"

    # ── Row 1: title headers ───────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(6)}1")
    ws["A1"] = "Achievable Vs Produced"
    ws["A1"].font = S["TITLE_L_FONT"]; ws["A1"].fill = S["TITLE_L_FILL"]
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    c1 = get_column_letter(CNTR_FIRST)
    c2 = get_column_letter(CNTR_FIRST + n_weeks - 1)
    ws.merge_cells(f"{c1}1:{c2}1")
    ws.cell(1, CNTR_FIRST, "By container")
    ws.cell(1, CNTR_FIRST).font = S["TITLE_C_FONT"]
    ws.cell(1, CNTR_FIRST).fill = S["TITLE_C_FILL"]
    ws.cell(1, CNTR_FIRST).alignment = Alignment(horizontal="center", vertical="center")

    p1 = get_column_letter(PCT_FIRST)
    p2 = get_column_letter(PCT_FIRST + n_weeks - 1)
    ws.merge_cells(f"{p1}1:{p2}1")
    ws.cell(1, PCT_FIRST, "By Percentage")
    ws.cell(1, PCT_FIRST).font = S["TITLE_R_FONT"]
    ws.cell(1, PCT_FIRST).fill = S["TITLE_R_FILL"]
    ws.cell(1, PCT_FIRST).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # ── Row 2: column headers ──────────────────────────────────────────────
    ws.row_dimensions[2].height = 36
    for c, h in {1:"Region", 2:None, 3:"Machine Line",
                  4:"No. of\nMachines", 5:"No. of\nShifts", 6:"Capacity(In\ncontainers)"}.items():
        cell = ws.cell(2, c, h)
        cell.font = S["HDR_FONT"]; cell.fill = S["HDR_FILL"]; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, lbl in enumerate(week_labels):
        for col, fill in [(CNTR_FIRST+i, S["CNTR_HDR_FILL"]),
                          (PCT_FIRST+i,  S["PCT_HDR_FILL"])]:
            cell = ws.cell(2, col, lbl)
            cell.font = S["HDR_FONT"]; cell.fill = fill; cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # ── Column widths ──────────────────────────────────────────────────────
    col_widths = {"A":12,"B":14,"C":28,"D":10,"E":8,"F":13}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w
    for i in range(n_weeks):
        ws.column_dimensions[get_column_letter(CNTR_FIRST+i)].width = 9.55 if i == 0 else 13
        ws.column_dimensions[get_column_letter(PCT_FIRST+i)].width  = 9.55 if i == 0 else 13

    # ── Hidden lookup table: cols after LAST_COL+2 ────────────────────────
    # Q-style cols: rate, shiftmin, tbgs_cfc, cfc_cntr per data row
    LKP_COL = LAST_COL + 2  # hidden start col
    hf = Font(name=FN, size=9, color="FFAAAAAA")
    for i, h in enumerate(["Rate", "ShiftMin", "TBGS/CFC", "CFC/Cntr"]):
        ws.cell(2, LKP_COL+i, h).font = hf
        ws.column_dimensions[get_column_letter(LKP_COL+i)].hidden = True

    # ── Data rows ──────────────────────────────────────────────────────────
    FIRST_DATA_ROW = 3
    row = FIRST_DATA_ROW
    pending_merges = []
    cur_a_start = cur_b_start = None
    cur_region = cur_sub = None
    grp_start = FIRST_DATA_ROW

    # Flatten layout to get data rows only (for capacity formula row tracking)
    data_rows = []   # list of actual sheet row numbers for data lines

    for item in LAYOUT:
        region, sub, line, nm, ns = item

        if region == "__TOTAL__":
            # TOTAL row
            is_total_row = row

            for c in range(1, LAST_COL+1):
                cell = ws.cell(row, c)
                cell.border = BORDER
                if c >= 3:
                    cell.fill = S["TOTAL_FILL"]
                    cell.font = S["BOLD"]

            ws.cell(row, 3, "TOTAL")
            ws.cell(row, 3).font = S["BOLD"]; ws.cell(row, 3).fill = S["TOTAL_FILL"]

            # F: SUM of group capacity
            f_col = get_column_letter(6)
            ws.cell(row, 6, f'=SUM(F{grp_start}:F{row-1})')
            ws.cell(row, 6).font = S["BOLD"]; ws.cell(row, 6).fill = S["TOTAL_FILL"]
            ws.cell(row, 6).number_format = S["NUMFMT_CNTR"]

            # G-J: SUM of achieved containers in group
            for i in range(n_weeks):
                col = get_column_letter(CNTR_FIRST+i)
                ach_range = f"{col}{grp_start}:{col}{row-1}"
                cap_range = f"F{grp_start}:F{row-1}"
                total_ach = sum(
                    achieved.get((ln, target_weeks[i]), 0)
                    for r2 in range(grp_start, row)
                    for ln in ([ws.cell(r2,3).value]
                               if ws.cell(r2,3).value not in ("TOTAL",None) else [])
                )
                total_cap = sum(
                    compute_capacity(ws.cell(r2,3).value,
                                     ws.cell(r2,4).value or 0,
                                     ws.cell(r2,5).value or 0)
                    for r2 in range(grp_start, row)
                    if ws.cell(r2,3).value not in ("TOTAL",None)
                )
                # Container TOTAL
                c_cell = ws.cell(row, CNTR_FIRST+i)
                c_cell.value = round(total_ach, 3) if total_ach else "-"
                c_cell.number_format = S["NUMFMT_CNTR"]
                c_cell.fill = S["TOTAL_FILL"]
                over = (total_cap > 0 and total_ach > total_cap)
                c_cell.font = S["OVER_FONT_B"] if over else S["BOLD"]
                c_cell.border = BORDER

                # Percentage TOTAL — weighted average formula
                p_cell = ws.cell(row, PCT_FIRST+i)
                c_str = get_column_letter(CNTR_FIRST+i)
                f_rng = f"F{grp_start}:F{row-1}"
                a_rng = f"{c_str}{grp_start}:{c_str}{row-1}"
                p_cell.value = (
                    f'=IFERROR(ROUND(SUMIF({f_rng},"<>-",{a_rng})'
                    f'/SUMIF({f_rng},"<>-",{f_rng})*100,0),0)'
                )
                p_cell.number_format = S["NUMFMT_PCT"]
                p_cell.fill = S["TOTAL_FILL"]
                # Over check
                pct_val = round(total_ach/total_cap*100) if total_cap else 0
                p_cell.font = S["OVER_FONT_B"] if pct_val > 100 else S["BOLD"]
                p_cell.border = BORDER

            ws.row_dimensions[row].height = 18
            grp_start = row + 1
            row += 1
            continue

        # ── Data line row ──────────────────────────────────────────────────
        ws.row_dimensions[row].height = 18

        # Region
        if region and region != cur_region:
            if cur_a_start is not None:
                pending_merges.append(("A", cur_a_start, row-1))
            cur_region = region; cur_a_start = row
            ws.cell(row, 1, region).font = S["REGION_FONT"]
            ws.cell(row, 1).fill = S["REGION_FILL"]

        # Sub-group
        if sub and sub != cur_sub:
            if cur_b_start is not None:
                pending_merges.append(("B", cur_b_start, row-1))
            cur_sub = sub; cur_b_start = row
            ws.cell(row, 2, sub).font = S["NORM"]

        ws.cell(row, 1).border = BORDER
        ws.cell(row, 2).border = BORDER

        # C — Machine Line
        ws.cell(row, 3, line).font = S["NORM"]; ws.cell(row, 3).border = BORDER

        # D — Machines (editable, yellow)
        ws.cell(row, 4, nm); ws.cell(row, 4).font = S["NORM"]
        ws.cell(row, 4).fill = S["INPUT_FILL"]; ws.cell(row, 4).border = BORDER
        ws.cell(row, 4).alignment = Alignment(horizontal="right")
        ws.cell(row, 4).number_format = "0"

        # E — Shifts (editable, yellow)
        ws.cell(row, 5, ns); ws.cell(row, 5).font = S["NORM"]
        ws.cell(row, 5).fill = S["INPUT_FILL"]; ws.cell(row, 5).border = BORDER
        ws.cell(row, 5).alignment = Alignment(horizontal="right")
        ws.cell(row, 5).number_format = "0"

        # Write hidden lookup factors
        f = MC_FACTORS.get(line, {})
        for i, key in enumerate(["rate","shiftmin","tbgs_cfc","cfc_cntr"]):
            ws.cell(row, LKP_COL+i, f.get(key, 0)).font = hf

        # F — Capacity: live formula using D, E, and hidden factors
        zero_cap = (ns == 0)
        f_cell = ws.cell(row, 6)
        f_cell.border = BORDER; f_cell.font = S["NORM"]
        f_cell.alignment = Alignment(horizontal="right")
        Q = get_column_letter(LKP_COL); R = get_column_letter(LKP_COL+1)
        S_col = get_column_letter(LKP_COL+2); T = get_column_letter(LKP_COL+3)
        if zero_cap:
            f_cell.value = "-"
        else:
            f_cell.value = (
                f'=IFERROR(IF({T}{row}=0,"-",'
                f'ROUND({Q}{row}*{R}{row}*$D{row}*$E{row}*6*0.9/{S_col}{row}/{T}{row},3)),'
                f'"-")'
            )
        f_cell.number_format = S["NUMFMT_CNTR"]

        # Pre-compute capacity for over-capacity colouring
        cap_val = compute_capacity(line, nm, ns)

        # G-J — Achieved containers (static from production register)
        for i, wn in enumerate(target_weeks):
            ach = achieved.get((line, wn), 0)
            c_cell = ws.cell(row, CNTR_FIRST+i)
            c_cell.border = BORDER
            c_cell.alignment = Alignment(horizontal="right")
            if zero_cap or (ach == 0 and not cap_val):
                c_cell.value = "-"; c_cell.font = S["NORM"]
            else:
                c_cell.value = round(ach, 3)
                c_cell.number_format = S["NUMFMT_CNTR"]
                c_cell.font = S["OVER_FONT"] if (cap_val and ach > cap_val) else S["NORM"]

        # K-N — Percentage: live formula referencing F (capacity)
        for i in range(n_weeks):
            ach_col = get_column_letter(CNTR_FIRST+i)
            p_cell = ws.cell(row, PCT_FIRST+i)
            p_cell.border = BORDER
            p_cell.alignment = Alignment(horizontal="right")
            if zero_cap:
                p_cell.value = "-"; p_cell.font = S["NORM"]
            else:
                p_cell.value = (
                    f'=IFERROR(IF($F{row}="-","-",'
                    f'IF($F{row}=0,0,ROUND({ach_col}{row}/$F{row}*100,0))),"-")'
                )
                p_cell.number_format = S["NUMFMT_PCT"]
                ach = achieved.get((line, target_weeks[i]), 0)
                p_cell.font = S["OVER_FONT"] if (cap_val and ach > cap_val) else S["NORM"]

        row += 1

    # ── Close merges ───────────────────────────────────────────────────────
    if cur_a_start is not None:
        pending_merges.append(("A", cur_a_start, row-1))
    if cur_b_start is not None:
        pending_merges.append(("B", cur_b_start, row-1))
    for cl, ds, de in pending_merges:
        if de > ds:
            ws.merge_cells(f"{cl}{ds}:{cl}{de}")
        ws[f"{cl}{ds}"].alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = f"{get_column_letter(CNTR_FIRST)}3"
    ws.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), skipped
